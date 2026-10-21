"""
Adds a "Warmstart vs. No Warmstart: Results So Far" slide (with the bar chart) to BOTH
srl_design_overview_en.pptx and figures_export/srl_meeting_deck.pptx.
"""
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x20, 0x20, 0x20)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(27); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(12.5); r2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top, width, font_size=13):
    tb = slide.shapes.add_textbox(Emu(600000), top, width, Emu(5000000))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "- ") + text
        r.font.name = "Calibri"; r.font.size = Pt(font_size - level * 1.5)
        r.font.color.rgb = BLACK if level == 0 else GRAY
        p.space_after = Pt(6 if level == 0 else 3)


BULLETS = [
    ("Same twin-critic config throughout (actor_lr=0.0032, critic_lr=0.0011, tau=0.005, "
     "gamma=0.99, reward_mode=local_positive, seed=42) - only use_actor_warmstart toggled", 0),
    ("20 epochs: warmstart val=0.6696/test=0.6841 vs. no-warmstart val=0.6443/test=0.6609 - "
     "only 2.5pp gap despite starting from scratch", 0),
    ("No-warmstart training is visibly less stable: val dipped to 0.1990 at epoch 10 (20-epoch "
     "run) and to 0.2134 at epoch 5 (30-epoch run) before recovering to 0.64+ by epoch 10-15", 0),
    ("30-epoch no-warmstart run in progress (2026-09-27) - checking whether it also approaches "
     "the warmstart 30-epoch result (val=0.7085)", 0),
    ("Source: local runs, srl_local_twin_critic_trial.py / srl_local_twin_critic_no_warmstart_"
     "trial.py / srl_local_twin_critic_epochs30_trial.py / srl_local_twin_critic_no_warmstart_"
     "epochs30_trial.py", 0),
]


def add_slide(deck_path):
    prs = Presentation(deck_path)
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)
    add_title(s, "Warmstart vs. No Warmstart: Results So Far",
              "Testing whether the reference paper's no-warmstart design costs much performance")
    add_bullets(s, BULLETS, top=Emu(1450000), width=Emu(11100000))
    s.shapes.add_picture("figures_export/SRL/warmstart_vs_no_warmstart_bar.png", Emu(2000000), Emu(3300000), width=Emu(7400000))
    prs.save(deck_path)
    print(f"Added slide to {deck_path} (now {len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    add_slide("srl_design_overview_en.pptx")
    add_slide("figures_export/srl_meeting_deck.pptx")
