"""
2026-10-14: adds one slide to reduce_then_optimize_deck.pptx (before "Summary & Open Points"): where the training
time of the c3 mixed-class runs goes (profile of the cluster logs, see
scripts/analysis/profile_log_phases.py and make_c3_mixedclass_runtime_by_file_chart.py) and the proposed lighter
training split "mixed-class-lite". Style copied from the neighbouring slides (Calibri, navy 132A4C titles, gray italic
5B6B7C subtitle, orange F2A541 badges, light cards).
Facts on the slide (all from the logs, seed 42):
  95-97 %       trip generation share of training time: 95.3 % (Baseline bi200/ss100) and 96.6 % (Pair-only bi400/ss80)
  87 % / 97 %   share of training time in lc203, lr203, lr204, lr206, lr207 (36.6 of 42.0 h, 140.5 of 144.7 h)
  25 min-3.3 h  average window time in these files (e.g. lc203: 18.9 h / 45 windows; lr207: 42.8 h / 13 windows)
                at only 7,000-30,000 feasible trips per window (max. per window: 6,987 .. 29,586)
Proposal numbers are ESTIMATES (8 h per epoch from the profile: ~4 h training + ~3.6 h validation).
Run: ./venv/bin/python3 add_c3_mixedclass_compute_slide.py
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

DECK = "reduce_then_optimize_deck.pptx"
FIG = "figures_export/Request Pruner SIL/Cardinality 3/c3_mixedclass_runtime_by_file.png"
NAVY = RGBColor(0x13, 0x2A, 0x4C)
GRAY = RGBColor(0x5B, 0x6B, 0x7C)
ORANGE = RGBColor(0xF2, 0xA5, 0x41)
INK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF3, 0xF5, 0xF7)
TITLE_TEXT = "c3 Mixed-Class: Where the Training Time Goes"

prs = Presentation(DECK)
if any(sh.has_text_frame and sh.text_frame.text.strip() == TITLE_TEXT for s in prs.slides for sh in s.shapes):
    raise SystemExit("slide already present - not adding it twice")
blank = prs.slide_layouts.get_by_name("Blank")
slide = prs.slides.add_slide(blank)


def text(left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    first = True
    for line in runs:  # each line = list of (text, size, bold, italic, color)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        for t, size, bold, italic, color in line:
            r = p.add_run()
            r.text = t
            r.font.name = "Calibri"; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.color.rgb = color
    return tb


# title + subtitle (same geometry as slide 29)
text(457200, 320040, 11247120, 640080, [[(TITLE_TEXT, 26, True, False, NAVY)]])
text(457200, 868680, 11247120, 457200,
     [[("Log profile of c3 mixed-class SIL runs (seed 42) - time per training file, and a proposed lighter training split",
        13, False, True, GRAY)]])

# chart (left)
pic_w = 6900000
pic = slide.shapes.add_picture(FIG, Emu(457200), Emu(1330000), width=Emu(pic_w))
pic_bottom = 1330000 + pic.height
text(457200, pic_bottom + 30000, pic_w, 300000,
     [[("Baseline: epoch 1 plus the start of epoch 2 (job cancelled); Pair-only: epoch 1 only. Both without Request Pruner.",
        11, False, True, GRAY)]])

# three stat callouts (right)
stats = [("95-97%", "of the training time is trip generation (RTV)"),
         ("87-97%", "of the training time sits in just 5 files: lc203, lr203, lr204, lr206, lr207"),
         ("25 min - 3.3 h", "average time per window in these files - at only 7,000-30,000 feasible trips (max. per window)")]
sx, sw = 7700000, 4050000
sy, sh_ = 1400000, 1250000
for k, (big, label) in enumerate(stats):
    y = sy + k * sh_
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(sx), Emu(y), Emu(sw), Emu(sh_ - 90000))
    card.adjustments[0] = 0.06
    card.fill.solid(); card.fill.fore_color.rgb = CARD
    card.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(sx), Emu(y), Emu(60000), Emu(sh_ - 90000))
    bar.fill.solid(); bar.fill.fore_color.rgb = ORANGE; bar.line.fill.background()
    text(sx + 200000, y + 60000, sw - 300000, 560000, [[(big, 32, True, False, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
    text(sx + 200000, y + 600000, sw - 300000, 500000, [[(label, 11.5, False, False, INK)]])

# proposal row (bottom)
py = pic_bottom + 420000
text(457200, py, 11247120, 300000, [[("Proposal: mixed-class-lite", 14, True, False, NAVY),
                                      ("   (estimates, decision pending)", 11, False, True, GRAY)]])
items = ["Train without lc203, lr203, lr204, lr206, lr207: 34 train files, the 12 val files stay unchanged",
         "Estimate: ~8 h per epoch at batch_interval=200/step_size=100, 5 epochs in ~2 days instead of > 9 days",
         "Trade-off: 4 of 7 lr2 train files removed; not directly comparable to c2 (39 files)"]
col_w, gap = 3600000, 224000
for k, item in enumerate(items):
    x = 457200 + k * (col_w + gap)
    y = py + 360000
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y + 20000), Emu(330000), Emu(330000))
    oval.fill.solid(); oval.fill.fore_color.rgb = ORANGE; oval.line.fill.background()
    tf = oval.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(k + 1)
    r.font.name = "Calibri"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE
    text(x + 430000, y, col_w - 430000, 620000, [[(item, 11, False, False, INK)]])

# move the new slide in front of the last slide ("Summary & Open Points")
lst = prs.slides._sldIdLst
ids = list(lst)
new = ids[-1]
lst.remove(new)
lst.insert(len(ids) - 2, new)
prs.save(DECK)
print("saved", DECK, "slides:", len(prs.slides))
