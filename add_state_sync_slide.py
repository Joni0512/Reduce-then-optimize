"""
2026-09-24: Adds a "Actor-RHO State Sync (Fixing Trajectory Divergence)" slide to BOTH
srl_design_overview_en.pptx and figures_export/srl_meeting_deck.pptx (see chat) - appended
at the end of each deck. Only adds this one slide per deck; touches nothing else.

Run: ./venv/bin/python3 add_state_sync_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x20, 0x20, 0x20)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(13); r2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top=Emu(1450000), font_size=13.5):
    tb = slide.shapes.add_textbox(Emu(600000), top, Emu(11100000), Emu(5600000))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "- ") + text
        r.font.name = "Calibri"; r.font.size = Pt(font_size - level * 1.5)
        r.font.color.rgb = BLACK if level == 0 else GRAY
        p.space_after = Pt(7 if level == 0 else 3)


BULLETS = [
    ("Problem: in the BC-vs-RHO setup (bi400/bi600/bi800), the actor runs its OWN "
     "independent iterative solve, only referencing RHO's (once-cached, fixed) manifest to "
     "reconstruct a y* imitation target at each step", 0),
    ("Once the actor's decision differs from RHO's even once, its vehicle state (active/"
     "picked-up requests, positions) DRIFTS away from RHO's - the two become two separate, "
     "diverging trajectories from a shared start", 0),
    ("This drift is the root cause behind ManifestConsistencyError / InfeasibleAssignmentError "
     "cases: RHO's manifest no longer matches the actor's actual (diverged) state", 0),
    ("Root cause confirmed in code: request visibility is a direct function of BATCH_INTERVAL "
     "(pickup window must overlap [current_time, current_time + BATCH_INTERVAL) - "
     "offline_rtv_solver.py/coaml_pipeline.py), not just the request's own arrival time - so a "
     "wider-horizon RHO structurally sees further ahead than the actor at the same current_time", 0),
    ("Fix (\"State Sync\"): after each actor iteration, overwrite the actor's vehicle state "
     "with RHO's own vehicle state at that same timestamp, filtered down to only the requests "
     "visible under the ACTOR's (smaller) BATCH_INTERVAL", 0),
    ("Implementation is cheap: vehicle state is not a persisted object - COAMLPipeline just "
     "threads a driver_runs dict list between iterations (coaml_pipeline.py ~line 380-381), "
     "rebuilding VehicleHandler/Vehicle from scratch every iteration - so injection only needs "
     "a filter helper + one conditional at that single line, no VehicleHandler changes", 0),
    ("Edge case: if RHO's vehicle is mid-route with a request not yet visible to the actor, "
     "the route is simply truncated to its visible portion (decided 2026-09-24 - accepted as a "
     "known simplification, not solved exactly)", 0),
    ("Expected effect: removes the actor/RHO state-divergence problem structurally, so "
     "ManifestConsistencyError/InfeasibleAssignmentError should become rare-to-absent, and the "
     "comparison becomes a cleaner test of whether a shorter-horizon actor can match RHO given "
     "an always-correct (projected) state, rather than being confounded by drift", 0),
    ("Status: design finalized 2026-09-24, not yet implemented - next: new "
     "srl_behavior_cloning_state_sync_loop.py + inject_vehicle_manifest_from_rho() helper", 0),
]


def add_slide(deck_path):
    prs = Presentation(deck_path)
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)
    add_title(s, "Actor-RHO State Sync: Fixing Trajectory Divergence",
              "Root cause of ManifestConsistencyError in BC-vs-RHO: independently-evolving actor state drifts from RHO's")
    add_bullets(s, BULLETS)
    prs.save(deck_path)
    print(f"Added slide to {deck_path} (now {len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    add_slide("srl_design_overview_en.pptx")
    add_slide("figures_export/srl_meeting_deck.pptx")
