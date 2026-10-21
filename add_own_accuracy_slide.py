"""
2026-09-23: Adds the new "Why bi200/ss100 Underperforms" diagnostic slide (own-accuracy
of each final prefix-label pruner, see analyze_prefix_pruner_own_accuracy.py) to
request_pruner_deck.pptx, right before "Summary & Open Points" (found by title text, not
a hardcoded index - the deck was externally re-saved once already and lost a slide, see
chat 2026-09-23). Only adds this one slide; touches nothing else.

Run: ./venv/bin/python3 add_own_accuracy_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

DECK = "request_pruner_deck.pptx"
SIL_DIR = "figures_export/Request Pruner SIL"

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(13.5); r2.font.color.rgb = GRAY


def add_caption(slide, text, top):
    tb = slide.shapes.add_textbox(Emu(457200), top, Emu(11277295), Emu(600000))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.name = "Calibri"; r.font.size = Pt(11); r.font.color.rgb = GRAY
    tb.text_frame.word_wrap = True


prs = Presentation(DECK)
BLANK = prs.slide_layouts[6]

insert_before = None
for i, s in enumerate(prs.slides):
    title = next((sh.text_frame.text for sh in s.shapes if sh.shape_type == 17), "")
    if title.strip() == "Summary & Open Points":
        insert_before = i
        break
if insert_before is None:
    raise SystemExit("Could not find 'Summary & Open Points' slide - aborting, nothing changed.")

s = prs.slides.add_slide(BLANK)
add_title(s, "Why Does bi200/ss100 Underperform? A Candidate Explanation",
          "Offline replay of each FINAL prefix-label pruner checkpoint — threshold=0.5, min_retention_fraction=0.3, "
          "class-1, window 0 (exact state, 59 vehicles/config)")
s.shapes.add_picture(f"{SIL_DIR}/prefix_pruner_own_accuracy_by_config.png",
                      Emu(2500000), Emu(1417320), Emu(7200000), Emu(6000000))
add_caption(s,
    "The pruner's own accuracy ranks exactly like the SIL outcomes: bi400/ss80 (best SIL result) loses its true "
    "target only 3.4% of the time; bi200/ss100 (worst SIL result) loses it 20.3% of the time — despite pruning "
    "almost the same overall share (42% both). The difference looks like PRECISION, not aggressiveness. "
    "Correlational across only 3 configs, not a proven cause — the window-count/divergence hypothesis tested earlier was refuted.",
    Emu(7550000))

xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[-1])
xml_slides.insert(insert_before, slides[-1])

prs.save(DECK)
print(f"saved {DECK}, {len(prs.slides)} slides, new slide inserted at index {insert_before}")
