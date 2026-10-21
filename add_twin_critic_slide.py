"""
2026-09-24: Adds a "Twin Critic (TD3-Style)" slide to BOTH srl_design_overview_en.pptx
and figures_export/srl_meeting_deck.pptx (see chat) - appended at the end of each deck
(newest development chronologically). Only adds this one slide per deck; touches nothing
else.

Run: ./venv/bin/python3 add_twin_critic_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x20, 0x20, 0x20)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(13.5); r2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top=Emu(1450000), font_size=14):
    tb = slide.shapes.add_textbox(Emu(600000), top, Emu(11100000), Emu(5500000))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "- ") + text
        r.font.name = "Calibri"; r.font.size = Pt(font_size - level * 1.5)
        r.font.color.rgb = BLACK if level == 0 else GRAY
        p.space_after = Pt(8 if level == 0 else 4)


BULLETS = [
    ("Motivation: earlier multi-seed diagnostic (critic vs. target_critic divergence, "
     "srl_meeting_deck slide 19) found unbounded Q-value overestimation drift under "
     "TD-bootstrap training", 0),
    ("TD3 (Fujimoto et al. 2018) fixes this with twin critics Q1/Q2: both trained against "
     "a SHARED min(Q1_target, Q2_target) bootstrap target; only Q1 used for the actor's "
     "policy gradient", 0),
    ("Our design (deliberate partial deviation from TD3 - see chat):", 0),
    ("TD-bootstrap target for training BOTH critics = min(target_critic, target_critic2)", 1),
    ("Actor's target-action ranking (softmax over candidate Q-values) = MEAN of "
     "target_critic and target_critic2, not just one - reduces ranking noise", 1),
    ("critic2: same architecture as critic (GAT, per TD3's own same-architecture-"
     "different-init convention - diversity comes from random init + training "
     "stochasticity, not structural differences)", 0),
    ("Same critic pretraining (10 epochs, full 38-instance train set), same Polyak "
     "target-update schedule/tau as the existing critic - fully symmetric setup", 0),
    ("Bug found + fixed while preparing this test: srl_training_loop.py's tau default was "
     "0.001, not the intended/validated 0.005 (run_srl_balanced_td_bootstrap_gat_"
     "bufferfix_12instances.py) - a copy-paste bug from 2026-09-13, affected all 22+ trials "
     "of the 38/9/9 actor_lr x critic_lr sweep so far; corrected for this test and going "
     "forward", 0),
    ("Status (2026-09-24): first single-critic-equivalent-config test running locally "
     "(seed=42, tau=0.005, gamma=0.99, full 38-train/9-val split, "
     "actor_lr=0.0032/critic_lr=0.0011 - best trial found in the sweep so far); a second, "
     "independent-seed replication (seed=43) queued on the cluster", 0),
]


def add_slide(deck_path, at_end=True):
    prs = Presentation(deck_path)
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)
    add_title(s, "Twin Critic (TD3-Style): Min-Target, Mean-Action-Selection",
              "Second independent GAT critic to fight Q-value overestimation drift - min for TD-bootstrap target, mean for actor's ranking")
    add_bullets(s, BULLETS)
    prs.save(deck_path)
    print(f"Added slide to {deck_path} (now {len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    add_slide("srl_design_overview_en.pptx")
    add_slide("figures_export/srl_meeting_deck.pptx")
