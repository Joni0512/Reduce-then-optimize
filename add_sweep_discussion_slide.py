"""
2026-09-25: Adds a compact "Sweep Results: Top Trials, Collapse Rate, Correlations" discussion
slide to BOTH srl_design_overview_en.pptx and figures_export/srl_meeting_deck.pptx (see chat) -
top 3 per reward_mode, mean of successful trials, crash/fail counts, actor_lr/critic_lr
correlations - NOT the full 33-row table (that's a separate export, see
figures_export/SRL/sweep_all_finished_runs_table.png/csv). Source: wandb project
srl-training-loop-sweep, pulled 2026-09-25 (159 runs total). tau=0.001 (known bug, see chat),
gamma=0.99, gnn_aggregator=gat, full 38-train/9-val split.

Run: ./venv/bin/python3 add_sweep_discussion_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x20, 0x20, 0x20)
GREEN = RGBColor(0x1B, 0x5E, 0x20)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(27); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(12.5); r2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top=Emu(1450000), font_size=14):
    tb = slide.shapes.add_textbox(Emu(600000), top, Emu(11100000), Emu(5500000))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "- ") + text
        r.font.name = "Calibri"; r.font.size = Pt(font_size - level * 1.5)
        r.font.color.rgb = BLACK if level == 0 else GRAY
        p.space_after = Pt(7 if level == 0 else 3)


BULLETS = [
    ("Source: wandb project srl-training-loop-sweep, pulled 2026-09-25 - 159 runs total: "
     "122 failed, 4 crashed, 33 finished. tau=0.001 (known bug, see chat), gamma=0.99, "
     "gnn_aggregator=gat, full 38-train/9-val split", 0),
    ("local (n=19 finished, 11 successful, 42% collapse rate): mean of successful = 0.6560", 0),
    ("Top 3: tough-sweep-42 (val=0.6892, actor_lr=2.19e-03, critic_lr=2.64e-03)", 1),
    ("eager-sweep-2 (val=0.6872, actor_lr=2.97e-04, critic_lr=2.09e-03)", 1),
    ("astral-sweep-48 (val=0.6848, actor_lr=2.03e-04, critic_lr=6.86e-03)", 1),
    ("local_positive (n=14 finished, 10 successful, 29% collapse rate): mean of successful = 0.4769", 0),
    ("Top 3: sandy-sweep-60 (val=0.6911, actor_lr=3.19e-03, critic_lr=1.08e-03)", 1),
    ("firm-sweep-68 (val=0.6826, actor_lr=9.59e-04, critic_lr=3.28e-03)", 1),
    ("glamorous-sweep-67 (val=0.6549, actor_lr=4.47e-04, critic_lr=6.40e-04)", 1),
    ("Correlations (log10 actor_lr/critic_lr vs. best_val_service_rate, n=33 finished, "
     "collapsed trials counted as val=0):", 0),
    ("corr(actor_lr, val) = -0.142 (weak)", 1),
    ("corr(critic_lr, val) = +0.436 (moderate - higher critic_lr tends to help)", 1),
    ("corr(critic_lr/actor_lr ratio, val) = +0.404 (moderate - critic learning faster than "
     "actor tends to help, consistent with the validated 12-instance default 10x ratio)", 1),
    ("Caveat: all of the above used the buggy tau=0.001, not the intended/validated tau=0.005 "
     "(see chat) - a corrected sweep is pending, current numbers are a reference point only", 0),
]


def add_slide(deck_path):
    prs = Presentation(deck_path)
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)
    add_title(s, "Sweep Results: Top Trials, Collapse Rate, Correlations",
              "srl-training-loop-sweep (actor_lr x critic_lr x reward_mode), 159 runs total")
    add_bullets(s, BULLETS)
    prs.save(deck_path)
    print(f"Added slide to {deck_path} (now {len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    add_slide("srl_design_overview_en.pptx")
    add_slide("figures_export/srl_meeting_deck.pptx")
