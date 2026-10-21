"""
2026-09-24: Restores the class-1 results table on the "Results for class 1 SIL Training
with the Expert-Prefix-Label Pruner" slide, lost in an external re-save (see chat
2026-09-24 - the deck lost this table plus the "Why bi200/ss100 Underperforms" slide a
second time). Reconstructed from build_request_pruner_deck.py's original table
(add_table call under "Slide 6: New results - service rate") plus the bi400/ss80 row
that update_request_pruner_deck_results.py had cloned in afterwards - WITH the
bi200/ss100 row's expert-prefix numbers corrected (60.9+-4.5%/13.6pp/61% -> 65.85+-2.45%/
8.65pp/75%, see chat 2026-09-24: the old numbers didn't match a fresh recompute from
sil_training_bi200_ss100_class1_requestpruner_prefixlabel_t0.5_seed{42,1,2,3,4}.log's
"Best val service rate" lines, confirmed independently by the user on the cluster).

Placed to the right of the existing chart picture (which was NOT lost and already
occupies the left ~48% of the slide at a repositioned, non-original location) rather
than below it, since below there is no longer room before the caption.

Run: ./venv/bin/python3 restore_class1_results_table.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DECK = "request_pruner_deck.pptx"
NAVY = RGBColor(0x1E, 0x27, 0x61)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADERS = ["Config", "No pruner", "RHO-label", "Expert-prefix", "Damage RHO / prefix", "Recovered"]
ROWS = [
    ["bi200/ss100", "74.5 ± 1.8%", "40.0 ± 5.9%", "65.85 ± 2.45%", "34.5 / 8.65pp", "75%"],
    ["bi400/ss200", "51.8 ± 1.6%", "43.6 ± 6.8%", "55.6 ± 1.7%", "8.2 / −3.8pp", "146% (beats baseline)"],
    ["bi400/ss80", "76.0 ± 1.0%", "57.9 ± 8.7%", "81.8 ± 0.6%", "18.1 / −5.8pp", "132% (beats baseline)"],
]
COL_WIDTHS = [1009000, 850000, 850000, 850000, 893000, 850000]

prs = Presentation(DECK)

target = None
for s in prs.slides:
    for sh in s.shapes:
        if sh.shape_type == 17 and sh.has_text_frame and sh.text_frame.text.startswith("Results for class 1 SIL Training"):
            target = s
            break
    if target:
        break
if target is None:
    raise SystemExit("Could not find the class-1 results slide - aborting, nothing changed.")

pic = next(sh for sh in target.shapes if sh.shape_type == 13)
left = pic.left + pic.width + Emu(200000)
top = pic.top
width = sum(COL_WIDTHS)
height = Emu(1800000)

n_rows, n_cols = len(ROWS) + 1, len(HEADERS)
shp = target.shapes.add_table(n_rows, n_cols, left, top, width, height)
tbl = shp.table
for c, w in zip(tbl.columns, COL_WIDTHS):
    c.width = Emu(w)

for j, h in enumerate(HEADERS):
    cell = tbl.cell(0, j)
    cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
    p = cell.text_frame.paragraphs[0]
    r = p.add_run(); r.text = h
    r.font.name = "Calibri"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

for i, row in enumerate(ROWS, start=1):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.name = "Calibri"; r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = NAVY
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT

prs.save(DECK)
print(f"saved {DECK}, table restored on '{target.shapes[0].text_frame.text[:40]}...' slide")
