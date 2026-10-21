"""
2026-09-25: Adds a small color-coded legend textbox to the "Refined: Pruning Sensitivity
Is Geometry-Dependent" slide, explaining what lc/lr/lrc mean directly on the slide (the
chart itself only has a compact "instance type" legend with no Li&Lim terminology spelled
out - user asked to see it "mit Farben" without needing chat context, 2026-09-25). Colors
match the chart's own palette (navy #1E2761 for lc, red #A6493A for lr/lrc).

Run: ./venv/bin/python3 add_lc_lr_color_legend.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

DECK = "request_pruner_deck.pptx"
NAVY_LC = RGBColor(0x1E, 0x27, 0x61)
RED_LR = RGBColor(0xA6, 0x49, 0x3A)
GRAY = RGBColor(0x60, 0x60, 0x60)

prs = Presentation(DECK)

target = None
for s in prs.slides:
    for sh in s.shapes:
        if sh.shape_type == 17 and sh.has_text_frame and sh.text_frame.text.startswith("Refined: Pruning Sensitivity Is Geometry-Dependent"):
            target = s
            break
    if target:
        break
if target is None:
    raise SystemExit("Could not find the 'Refined: Pruning Sensitivity...' slide - aborting, nothing changed.")

tb = target.shapes.add_textbox(Emu(457200), Emu(1000000), Emu(2000000), Emu(1200000))
tf = tb.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
r1a = p1.add_run(); r1a.text = "● lc"
r1a.font.name = "Calibri"; r1a.font.size = Pt(12); r1a.font.bold = True; r1a.font.color.rgb = NAVY_LC
r1b = p1.add_run(); r1b.text = " = clustered\npickup/dropoff\nlocations"
r1b.font.name = "Calibri"; r1b.font.size = Pt(10.5); r1b.font.color.rgb = GRAY

p2 = tf.add_paragraph()
p2.space_before = Pt(8)
r2a = p2.add_run(); r2a.text = "● lr / lrc"
r2a.font.name = "Calibri"; r2a.font.size = Pt(12); r2a.font.bold = True; r2a.font.color.rgb = RED_LR
r2b = p2.add_run(); r2b.text = " = random /\nmixed locations"
r2b.font.name = "Calibri"; r2b.font.size = Pt(10.5); r2b.font.color.rgb = GRAY

prs.save(DECK)
print(f"saved {DECK}")
