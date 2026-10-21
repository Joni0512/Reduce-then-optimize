"""
2026-09-24: Adds the cross-config pruning-rate-vs-service-rate slide (all 3 configs) -
refines/replaces the single-config (bi200/ss100 only) framing from
add_pruning_rate_scatter_slide.py with the more robust interaction finding: same
pruning rate hurts lr/lrc instances far more than lc instances, consistent across all
3 configs (a pure "pruning rate correlates with service rate" story does NOT hold
uniformly - bi400/ss200's lr111/lr112/lrc107/lrc108 are barely pruned at all yet still
collapse). Inserted right after the existing single-config scatter slide.

Run: ./venv/bin/python3 add_pruning_interaction_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

DECK = "request_pruner_deck.pptx"
SIL_DIR = "figures_export/Request Pruner SIL"

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(13); r2.font.color.rgb = GRAY


def add_caption(slide, text, top):
    tb = slide.shapes.add_textbox(Emu(457200), top, Emu(11277295), Emu(700000))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.name = "Calibri"; r.font.size = Pt(11); r.font.color.rgb = GRAY
    tb.text_frame.word_wrap = True


prs = Presentation(DECK)
BLANK = prs.slide_layouts[6]

insert_after = None
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.shape_type == 17 and sh.has_text_frame and sh.text_frame.text.startswith("Root Cause of the Instance Gap: Pruning Rate"):
            insert_after = i
            break
    if insert_after is not None:
        break
if insert_after is None:
    raise SystemExit("Could not find the single-config pruning-rate slide - aborting, nothing changed.")

s = prs.slides.add_slide(BLANK)
add_title(s, "Refined: Pruning Sensitivity Is Geometry-Dependent",
           "All 3 configs, mixed-class, expert-prefix-label pruner - seed 42 pruning rate vs. 5-seed-mean service rate")
s.shapes.add_picture(f"{SIL_DIR}/pruning_vs_service_rate_all_configs_scatter.png",
                      Emu(2438400), Emu(1264920), Emu(7315200))
add_caption(s,
    "A simple 'more pruning = less service' story only holds within bi200/ss100. Across all 3 configs, the "
    "robust pattern is an interaction: lc instances (clustered) tolerate even 23-51% pruning while staying "
    "85-95% served; lr/lrc instances (random/mixed) collapse toward 0-36% at the SAME or even lower pruning "
    "rates - bi400/ss200's lr111/lr112/lrc107/lrc108 are barely pruned at all (0-4%) yet still only reach "
    "32-36%. lc instances are pruning-robust (redundant, spatially clustered trip alternatives survive "
    "pruning); lr/lrc instances are pruning-fragile (few good candidates to begin with).",
    Emu(7700000))

xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[-1])
xml_slides.insert(insert_after + 1, slides[-1])

prs.save(DECK)
print(f"saved {DECK}, {len(prs.slides)} slides, new slide inserted at index {insert_after + 1}")
