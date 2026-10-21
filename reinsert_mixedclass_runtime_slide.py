"""
2026-09-23: Re-inserts the "Mixed-Class Runtime & Savings" slide that update_request_pruner_deck_v2.py
added, but which was lost when the deck was re-saved externally afterward (mtime moved past my
save, slide count back to 9, slide 6 also lost its title textbox - see chat). This script only
adds the missing slide back at the correct position (after "New Results: Mixed-Class", before
"Summary & Open Points") - it does NOT touch anything else, so whatever the user changed
externally in the meantime is left alone.

Run: ./venv/bin/python3 reinsert_mixedclass_runtime_slide.py
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

DECK = "request_pruner_deck.pptx"
SIL_DIR = "figures_export/Request Pruner SIL"

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(14); r2.font.color.rgb = GRAY


def add_caption(slide, text, top):
    tb = slide.shapes.add_textbox(Emu(457200), top, Emu(11277295), Emu(457200))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.name = "Calibri"; r.font.size = Pt(11); r.font.color.rgb = GRAY
    tb.text_frame.word_wrap = True


prs = Presentation(DECK)
BLANK = prs.slide_layouts[6]

# locate current index of "New Results: Mixed-Class" (service rate slide) to insert right after it
insert_after = None
for i, s in enumerate(prs.slides):
    title = next((sh.text_frame.text for sh in s.shapes if sh.shape_type == 17), "")
    if title.strip() == "New Results: Mixed-Class":
        insert_after = i
        break
if insert_after is None:
    raise SystemExit("Could not find 'New Results: Mixed-Class' slide - aborting, nothing changed.")

s = prs.slides.add_slide(BLANK)
add_title(s, "New Results: Mixed-Class Runtime & Savings",
          "Same runs as the previous slide — total wall-clock (train+val), % = saving vs. no-pruner baseline")
s.shapes.add_picture(f"{SIL_DIR}/sil_mixedclass_requestpruner_prefixlabel_final_comparison_runtime.png",
                      Emu(2895600), Emu(1417320), Emu(6400800), Emu(6600000))
add_caption(s, "The expert-prefix-label pruner saves about as much runtime as the RHO-label pruner at bi200/ss100 (63.6% vs. 62.9%) "
               "and somewhat less at bi400/ss200 (34.1% vs. 39.9%), while beating the baseline's service rate there.", Emu(6800000))

xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[-1])
xml_slides.insert(insert_after + 1, slides[-1])

prs.save(DECK)
print(f"saved {DECK}, {len(prs.slides)} slides, new slide inserted at index {insert_after + 1}")
