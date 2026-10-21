"""
Adds a "SIL Warmstart: Design Deviation + Data Leakage Finding" slide to BOTH
srl_design_overview_en.pptx and figures_export/srl_meeting_deck.pptx (see chat).
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
GRAY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x20, 0x20, 0x20)


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Emu(457200), Emu(292608), Emu(11277295), Emu(640080))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.name = "Cambria"; r.font.size = Pt(27); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Emu(457200), Emu(932688), Emu(11277295), Emu(365760))
        p2 = tb2.text_frame.paragraphs[0]
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.name = "Calibri"; r2.font.size = Pt(12.5); r2.font.color.rgb = GRAY


def add_bullets(slide, bullets, top=Emu(1450000), font_size=14):
    tb = slide.shapes.add_textbox(Emu(600000), top, Emu(11100000), Emu(5500000))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "- ") + text
        r.font.name = "Calibri"; r.font.size = Pt(font_size - level * 1.5)
        r.font.color.rgb = BLACK if level == 0 else GRAY
        p.space_after = Pt(7 if level == 0 else 3)


BULLETS = [
    ("Our setup: every twin-critic run so far warm-starts the actor from a SIL checkpoint "
     "(sil_training_bi200_ss100_mixed_balanced_legacy_mlp_seed1) before RL training begins", 0),
    ("Reference paper check (Hoppe et al., \"Structured Reinforcement Learning for "
     "Combinatorial Decision-Making\", NeurIPS 2025) - Section C.1.3:", 0),
    ("\"we use the same random model initialization for SIL, PPO, and SRL\"", 1),
    ("SRL is trained purely from interaction, no SIL warmstart - SIL and SRL are independent "
     "baselines, not a sequential pretrain-then-finetune pipeline", 1),
    ("=> our warmstart design is a deliberate deviation from the reference paper, not a "
     "direct reproduction - must be stated explicitly in the thesis", 0),
    ("Data leakage finding while investigating this (2026-09-26): the SIL checkpoint's own "
     "train/val split OVERLAPS with our 38/9/9 SRL split", 0),
    ("4 of 9 VAL_INSTANCES (lc202, lrc102, lrc202, lr202) were in the checkpoint's SIL "
     "TRAINING set", 1),
    ("4 of 9 TEST_INSTANCES (lc208, lr111, lrc108, lc108) were in the checkpoint's SIL "
     "VALIDATION set", 1),
    ("=> every twin-critic val/test service rate reported so far may be optimistically "
     "biased by this leakage", 1),
    ("New experiment (2026-09-26): use_actor_warmstart=False flag added to "
     "run_srl_training_loop() - actor stays at its own random init throughout (same one "
     "reused across critic pretraining + main training, not re-randomized per call)", 0),
    ("Running: twin-critic, no warmstart, same LRs/tau/architecture as our best run, "
     "multiple seeds (local seed 42 + cluster seeds 43/44/45) - results pending", 0),
]


def add_slide(deck_path):
    prs = Presentation(deck_path)
    BLANK = prs.slide_layouts[6]
    s = prs.slides.add_slide(BLANK)
    add_title(s, "SIL Warmstart: Design Deviation + Data Leakage Finding",
              "Reference-paper comparison uncovered both a methodological deviation and a real train/val/test leak")
    add_bullets(s, BULLETS)
    prs.save(deck_path)
    print(f"Added slide to {deck_path} (now {len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    add_slide("srl_design_overview_en.pptx")
    add_slide("figures_export/srl_meeting_deck.pptx")
